from __future__ import annotations

import csv
from datetime import UTC, datetime
from uuid import UUID

import pytest
from openpyxl import load_workbook

from fb_crawl.contacts.models import LookupOutcome, LookupSource
from fb_crawl.exports.formatters import (
    EXPORT_COLUMNS,
    safe_cell,
    write_history_csv,
    write_history_xlsx,
)
from fb_crawl.history.models import HistoryItem


NOW = datetime(2026, 8, 31, 3, tzinfo=UTC)


def item(*, name: str = "Sample User", phone: str = "+84981234567"):
    return HistoryItem(
        id=71,
        account_id=7,
        device_id=9,
        facebook_user_id=41,
        facebook_uid="100123",
        username="sample.user",
        name=name,
        profile_url="https://www.facebook.com/sample.user",
        phone_number_id=81,
        phone=phone,
        outcome=LookupOutcome.FOUND,
        source=LookupSource.CACHE,
        provider_called=False,
        quota_charged=True,
        safe_error_code="",
        created_at=NOW,
        completed_at=NOW,
        scan_mode="manual_loaded",
        source_type="post_author",
        source_url="https://www.facebook.com/groups/123/posts/456",
        product_crawl_job_id=UUID(
            "11111111-1111-4111-8111-111111111111"
        ),
    )


@pytest.mark.parametrize("value", ["=1+1", "+cmd", "-2", "@SUM(A1:A2)"])
def test_spreadsheet_formula_values_are_neutralized(value: str) -> None:
    assert safe_cell(value).startswith("'")


def test_csv_is_utf8_and_neutralizes_every_user_controlled_cell(tmp_path) -> None:
    path = tmp_path / "history.csv"

    write_history_csv((item(name="=HYPERLINK('x')"),), path)

    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert tuple(rows[0]) == EXPORT_COLUMNS
    assert EXPORT_COLUMNS == (
        "event_id",
        "facebook_uid",
        "username",
        "name",
        "profile_url",
        "phone",
        "gender",
        "address",
        "birth_date",
        "outcome",
        "source",
        "scan_mode",
        "source_type",
        "source_url",
        "product_crawl_job_id",
        "provider_called",
        "quota_charged",
        "created_at",
        "completed_at",
    )
    assert rows[0]["scan_mode"] == "manual_loaded"
    assert rows[0]["source_type"] == "post_author"
    assert rows[0]["source_url"] == (
        "https://www.facebook.com/groups/123/posts/456"
    )
    assert rows[0]["name"].startswith("'")
    assert rows[0]["phone"].startswith("'")


def test_xlsx_uses_write_only_output_and_safe_cells(tmp_path) -> None:
    path = tmp_path / "history.xlsx"

    write_history_xlsx((item(name="@SUM(A1:A2)"),), path)

    workbook = load_workbook(path, read_only=True, data_only=False)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0] == EXPORT_COLUMNS
    assert rows[1][EXPORT_COLUMNS.index("name")].startswith("'")
    assert rows[1][EXPORT_COLUMNS.index("phone")].startswith("'")
